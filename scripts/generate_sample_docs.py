"""Generate fictional sample company documents (PDF, DOCX, TXT) for demos and tests.

Usage:
    python scripts/generate_sample_docs.py              # writes to sample_documents/
    python scripts/generate_sample_docs.py --install    # also copies them into knowledge_base/

The documents intentionally contain two conflicts for testing conflict detection:
  * Remote work days: Employee Handbook says 3 days/week, Remote Work Policy says 2 days/week.
  * Meal allowance: Employee Handbook says $50/day, Expense Reimbursement SOP says $75/day.
"""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sample_documents"
COMPANY = "Brightpath Solutions"

HANDBOOK = [
    ("1. Welcome", [
        f"This Employee Handbook describes the policies that apply to all employees of {COMPANY}. "
        "Effective date: 1 January 2025. Owner: People Operations.",
    ]),
    ("2. Working Hours", [
        "Core working hours are 10:00 to 16:00 local time, Monday to Friday. Full-time employees work 40 hours per week.",
        "Overtime must be approved in advance by your line manager.",
    ]),
    ("3. Paid Time Off", [
        "Full-time employees receive 20 days of paid time off (PTO) per calendar year, accrued monthly.",
        "Up to 5 unused PTO days may be carried over into the next calendar year. Carried-over days expire on 31 March.",
        "PTO requests of 3 days or more must be submitted at least 2 weeks in advance through the HR portal.",
        "Employees also receive 10 paid sick days per year. A doctor's note is required for absences longer than 3 consecutive days.",
    ]),
    ("4. Remote Work", [
        "Employees may work remotely up to 3 days per week with their manager's approval.",
        "Remote employees must be reachable during core working hours and use the company VPN.",
    ]),
    ("5. Travel and Meals", [
        "When travelling on company business, the meal allowance is $50 per day.",
        "Economy class is required for flights under 6 hours.",
    ]),
    ("6. Probation Period", [
        "New employees serve a 90-day probation period. A review meeting is held with the line manager at the end of probation.",
    ]),
    ("7. Code of Conduct", [
        "Employees must treat colleagues, customers and partners with respect. Harassment and discrimination are not tolerated.",
        "Report concerns to People Operations or anonymously via the Ethics Hotline. Retaliation against anyone who raises a concern in good faith is prohibited.",
    ]),
    ("8. Information Security", [
        "Never share passwords. Enable multi-factor authentication on all company accounts.",
        "Report lost or stolen devices to the IT Service Desk within 1 hour.",
    ]),
]

ONBOARDING = [
    ("NEW HIRE ONBOARDING GUIDE", [f"This guide helps new {COMPANY} employees through their first 30 days. Owner: People Operations."]),
    ("Before Day One", [
        "1. Sign your employment contract and return it via the HR portal.",
        "2. Complete the background check form.",
        "3. IT ships your laptop to arrive at least 2 business days before your start date.",
    ]),
    ("Week One", [
        "1. Attend the company orientation session on Monday at 09:30.",
        "2. Set up multi-factor authentication and the company VPN with the IT Service Desk.",
        "3. Meet your onboarding buddy, who is assigned by your manager.",
        "4. Complete mandatory Information Security training within 5 business days.",
    ]),
    ("First 30 Days", [
        "1. Agree on 30-60-90 day goals with your manager by the end of week two.",
        "2. Complete Code of Conduct training within 30 days.",
        "3. Schedule a 30-day check-in with People Operations.",
    ]),
]

REMOTE_POLICY = [
    ("Remote Work Policy", "title", [f"{COMPANY} - Effective date: 1 July 2026. Owner: People Operations."]),
    ("Eligibility", "h1", [
        "Employees who have completed their probation period are eligible for hybrid remote work.",
        "Employees may work remotely up to 2 days per week. Team anchor days (Tuesday and Thursday) are in the office.",
    ]),
    ("Approval Process", "h1", [
        "Submit a remote work request in the HR portal.",
        "Your line manager approves or rejects the request within 5 business days.",
        "Fully remote arrangements require approval from the department head and People Operations.",
    ]),
    ("Equipment and Stipend", "h1", [
        "Remote-eligible employees receive a one-time home office stipend of $500.",
        "Company laptops must be used for all work. Personal devices may not store company data.",
    ]),
    ("Security Requirements", "h1", [
        "Always connect through the company VPN. Do not work from public Wi-Fi without the VPN.",
        "Lock your screen whenever you step away from your device.",
    ]),
]

VENDOR_PROCEDURE = [
    ("Vendor Management Procedure", "title", [f"{COMPANY} Procurement. Applies to all purchases of goods and services."]),
    ("1. Vendor Selection", "h1", [
        "For purchases over $10,000, obtain at least three written quotes from different vendors.",
        "Purchases under $10,000 require one quote and line manager approval.",
    ]),
    ("2. Vendor Onboarding", "h1", [
        "Request the vendor's W-9 or tax registration and banking details through the Procurement portal.",
        "Vendors that process company or customer data must pass an Information Security review before contract signature.",
        "Legal reviews all contracts with a value above $25,000.",
    ]),
    ("3. Performance Reviews", "h1", [
        "Strategic vendors are reviewed quarterly against agreed service levels.",
        "Underperforming vendors are placed on a 60-day improvement plan.",
    ]),
]
VENDOR_TABLE = [
    ("Purchase value", "Quotes required", "Approver"),
    ("Under $10,000", "1", "Line manager"),
    ("$10,000 - $50,000", "3", "Department head"),
    ("Over $50,000", "3", "CFO"),
]

EXPENSE_SOP = f"""EXPENSE REIMBURSEMENT SOP
{COMPANY} Finance - Standard Operating Procedure FIN-004. Effective date: 1 March 2026. Owner: Finance.

PURPOSE
This procedure explains how employees claim reimbursement for business expenses.

ELIGIBLE EXPENSES
Business travel, client meals, conference fees and pre-approved training are eligible.
The meal allowance during business travel is $75 per day.
Alcohol is not reimbursable unless pre-approved for client entertainment.

SUBMISSION STEPS
1. Keep an itemised receipt for every expense over $25.
2. Log in to the ExpenseHub portal and select "New claim".
3. Attach receipts and choose the correct cost centre.
4. Submit the claim within 30 days of the expense date.
5. Your line manager approves the claim within 5 business days.
6. Finance reimburses approved claims in the next payroll run.

LATE OR MISSING RECEIPTS
Claims submitted after 30 days require Finance Director approval.
If a receipt is lost, complete the Missing Receipt Declaration form.
"""

SUPPORT_GUIDELINES = f"""CUSTOMER SUPPORT GUIDELINES
{COMPANY} Customer Support Team. Owner: Head of Customer Support.

TICKET PRIORITIES
P1 - Critical: service is down for all users. First response within 1 hour, 24/7.
P2 - High: major feature broken with no workaround. First response within 4 business hours.
P3 - Normal: issue with a workaround or a general question. First response within 1 business day.

ESCALATION PROCESS
1. If a P1 ticket is not acknowledged within 30 minutes, page the on-call support engineer.
2. If a P1 is not resolved within 4 hours, escalate to the Head of Customer Support.
3. The Head of Customer Support notifies the customer's account manager and posts updates every 60 minutes.

COMMUNICATION STANDARDS
Address customers by name and use a friendly, professional tone.
Never share internal ticket notes or other customers' information.
Close a ticket only after the customer confirms the issue is resolved, or after 5 business days without a reply.

REFUNDS
Support agents may approve refunds up to $200. Larger refunds require Support Manager approval.
"""


def make_pdf(path: Path, title: str, sections, page_break_every: int = 3) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for i, (heading, paras) in enumerate(sections):
        if i and i % page_break_every == 0:
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story.append(Paragraph(p, styles["BodyText"]))
            story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


def make_docx(path: Path, sections, table=None) -> None:
    import docx

    document = docx.Document()
    for heading, kind, paras in sections:
        if kind == "title":
            document.add_heading(heading, level=0)
        else:
            document.add_heading(heading, level=1)
        for p in paras:
            document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.save(str(path))


def generate(out_dir: Path = OUT) -> list:
    out_dir.mkdir(parents=True, exist_ok=True)
    make_pdf(out_dir / "employee_handbook.pdf", f"{COMPANY} Employee Handbook", HANDBOOK)
    make_pdf(out_dir / "onboarding_guide.pdf", f"{COMPANY} Onboarding Guide", ONBOARDING, page_break_every=10)
    make_docx(out_dir / "remote_work_policy.docx", REMOTE_POLICY)
    make_docx(out_dir / "vendor_management_procedure.docx", VENDOR_PROCEDURE, VENDOR_TABLE)
    (out_dir / "expense_reimbursement_sop.txt").write_text(EXPENSE_SOP, encoding="utf-8")
    (out_dir / "customer_support_guidelines.txt").write_text(SUPPORT_GUIDELINES, encoding="utf-8")
    return sorted(out_dir.iterdir())


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--install", action="store_true", help="Copy samples into knowledge_base/")
    args = parser.parse_args()
    files = generate()
    print("Generated:", *[f.name for f in files], sep="\n  ")
    if args.install:
        kb = ROOT / "knowledge_base"
        kb.mkdir(exist_ok=True)
        for f in files:
            shutil.copy2(f, kb / f.name)
        print(f"Copied {len(files)} files into {kb}")
